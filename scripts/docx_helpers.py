#!/usr/bin/env python3
"""
Script per generare il documento Word professionale:
'Guida_Costruzione_Prompt_QSA_CounselorBot.docx'

Utilizza:
- /home/nugh75/counselorbot-sbs/docs/qsa_prompt_strategie_letture_db.txt
- /home/nugh75/counselorbot-sbs/docs/qsa_prompt_strategie_letture.txt
- Codice di backend per il prompt pipeline (chat_preparation.py, chat_logic.py)
"""

import os
import sys
import json
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colori Palette Corporate Professionale
NAVY = RGBColor(26, 54, 93)       # #1A365D - Titoli principali e intestazioni
TEAL = RGBColor(13, 148, 136)     # #0D9488 - Sottotitoli e accenti teorici
SLATE = RGBColor(51, 65, 85)      # #334155 - Corpo testo primario
CHARCOAL = RGBColor(30, 41, 59)   # #1E293B - Testo scuro
MUTED = RGBColor(100, 116, 139)   # #64748B - Metadati e note a margine
AMBER = RGBColor(180, 83, 9)      # #B45309 - Warning / Vincoli critici

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, title, text, border_color="2563EB", bg_color="F8FAFC", title_color="1E3A8A", font_name="Consolas", font_size=9.0):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor.from_string(title_color)

    run_body = p.add_run(text)
    run_body.font.name = font_name
    run_body.font.size = Pt(font_size)
    run_body.font.color.rgb = RGBColor(51, 65, 85)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_theory_box(doc, title, text):
    add_callout_box(
        doc,
        title=f"FRAMEWORK TEORICO PELLEREY: {title}",
        text=text,
        border_color="0D9488",
        bg_color="F0FDFA",
        title_color="0F766E",
        font_name="Calibri",
        font_size=9.5
    )

def add_prompt_box(doc, title, text):
    add_callout_box(
        doc,
        title=f"PROMPT DAL DATABASE (configs/guided_steps): {title}",
        text=text,
        border_color="2563EB",
        bg_color="F8FAFC",
        title_color="1D4ED8",
        font_name="Consolas",
        font_size=9.0
    )

def add_warning_box(doc, title, text):
    add_callout_box(
        doc,
        title=f"VINCOLO DI SISTEMA E REGOLE DI INVERSIONE: {title}",
        text=text,
        border_color="D97706",
        bg_color="FFFBEB",
        title_color="B45309",
        font_name="Calibri",
        font_size=9.5
    )

def add_heading_1(doc, text):
    h = doc.add_heading(level=1)
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(level=2)
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = TEAL
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    return h

def add_heading_3(doc, text):
    h = doc.add_heading(level=3)
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = CHARCOAL
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    return h

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = CHARCOAL
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = SLATE
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = CHARCOAL
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = SLATE
    return p

def add_styled_table(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)

    # Header Row
    hdr_row = tbl.rows[0]
    set_row_header(hdr_row)
    set_row_cant_split(hdr_row)
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(heading)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_data in enumerate(rows_data):
        row = tbl.rows[r_idx + 1]
        set_row_cant_split(row)
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = SLATE

    # Column widths
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return tbl

print("Modulo helper inizializzato.")
